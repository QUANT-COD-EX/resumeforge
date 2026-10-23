import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_resume_docx(resume_data: dict) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.65)
        s.right_margin = Inches(0.65)

    personal = resume_data.get("personal", {})
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_title.add_run(personal.get("fullName") or "Your Name")
    r_name.bold = True
    r_name.font.size = Pt(18)

    contact = " | ".join(filter(None, [
        personal.get("title"), personal.get("email"),
        personal.get("phone"), personal.get("location")
    ]))
    if contact:
        p_sub = doc.add_paragraph(contact)
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if personal.get("summary"):
        doc.add_heading("Professional Summary", level=2)
        doc.add_paragraph(personal["summary"])

    for exp in resume_data.get("experience", []):
        doc.add_heading(f"{exp.get('role', '')} — {exp.get('company', '')}", level=3)
        for b in exp.get("bullets", []):
            if b.strip():
                doc.add_paragraph(b.strip(), style="List Bullet")

    for proj in resume_data.get("projects", []):
        doc.add_heading(f"{proj.get('name', '')} ({proj.get('techStack', '')})", level=3)
        for b in proj.get("bullets", []):
            if b.strip():
                doc.add_paragraph(b.strip(), style="List Bullet")

    for edu in resume_data.get("education", []):
        doc.add_heading("Education", level=2)
        doc.add_paragraph(f"{edu.get('degree', '')}, {edu.get('institution', '')} ({edu.get('startDate', '')} - {edu.get('endDate', '')})")

    skills = resume_data.get("skills", {})
    all_skills = skills.get("technical", []) + skills.get("tools", []) + skills.get("soft", [])
    if all_skills:
        doc.add_heading("Skills", level=2)
        doc.add_paragraph(", ".join(all_skills))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
